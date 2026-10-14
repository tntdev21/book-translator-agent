#!/usr/bin/env python3
"""Extract text from DOCX to Markdown"""

import sys
import os

def extract_docx(input_path, output_path):
    """
    Extract text from DOCX and save as Markdown.
    Preserves headings, paragraphs, and basic formatting.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    except ImportError:
        print("❌ Missing python-docx. Run: pip install python-docx --break-system-packages")
        sys.exit(1)
    
    print(f"📄 Extracting from {input_path}...")
    
    doc = Document(input_path)
    content = []
    
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        
        if not text:
            continue
        
        # Detect headings by style
        style_name = para.style.name.lower() if para.style else ""
        
        if 'heading 1' in style_name or 'title' in style_name:
            content.append(f"# {text}\n")
        elif 'heading 2' in style_name:
            content.append(f"## {text}\n")
        elif 'heading 3' in style_name:
            content.append(f"### {text}\n")
        elif 'heading 4' in style_name:
            content.append(f"#### {text}\n")
        else:
            # Regular paragraph
            # Check for bold/italic (simplified - checks first run)
            if para.runs:
                first_run = para.runs[0]
                if first_run.bold and len(para.runs) == 1:
                    content.append(f"**{text}**\n")
                elif first_run.italic and len(para.runs) == 1:
                    content.append(f"*{text}*\n")
                else:
                    content.append(f"{text}\n")
            else:
                content.append(f"{text}\n")
        
        # Progress
        if (i + 1) % 100 == 0:
            print(f"  ✓ Processed {i+1} paragraphs...")
    
    # Extract tables
    for table in doc.tables:
        content.append("\n")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            content.append("| " + " | ".join(cells) + " |")
        content.append("\n")
    
    # Write output
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(content))
    
    total_words = sum(len(p.split()) for p in content)
    print(f"✅ Done! Saved to {output_path}")
    print(f"   Paragraphs: {len(doc.paragraphs)}")
    print(f"   Tables: {len(doc.tables)}")
    print(f"   Words: ~{total_words:,}")

def extract_docx_by_pages(input_path, output_path, max_pages=None):
    """
    Extract DOCX with page markers (approximate).
    DOCX doesn't have real page info, so we estimate ~500 words per page.
    """
    try:
        from docx import Document
    except ImportError:
        print("❌ Missing python-docx. Run: pip install python-docx --break-system-packages")
        sys.exit(1)
    
    print(f"📄 Extracting from {input_path}...")
    
    doc = Document(input_path)
    content = []
    
    words_per_page = 500
    current_words = 0
    current_page = 1
    content.append(f"<!-- Page {current_page} -->\n")
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        words = len(text.split())
        current_words += words
        
        # Check for new page (approximate)
        if current_words >= words_per_page:
            current_page += 1
            if max_pages and current_page > max_pages:
                break
            content.append(f"\n<!-- Page {current_page} -->\n")
            current_words = 0
        
        # Detect headings
        style_name = para.style.name.lower() if para.style else ""
        
        if 'heading 1' in style_name or 'title' in style_name:
            content.append(f"# {text}\n")
        elif 'heading 2' in style_name:
            content.append(f"## {text}\n")
        elif 'heading 3' in style_name:
            content.append(f"### {text}\n")
        else:
            content.append(f"{text}\n")
    
    # Write output
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(content))
    
    print(f"✅ Done! Saved to {output_path}")
    print(f"   Estimated pages: {current_page}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python extract_docx.py <input.docx> [output.md] [max_pages]")
        print("Example: python extract_docx.py input/source.docx source/full_text.md")
        print("Example: python extract_docx.py input/sample.docx source/sample.md 30")
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "source/full_text.md"
    max_pages = int(sys.argv[3]) if len(sys.argv) > 3 else None
    
    if not os.path.exists(input_file):
        print(f"❌ Error: File not found: {input_file}")
        sys.exit(1)
    
    if max_pages:
        extract_docx_by_pages(input_file, output_file, max_pages)
    else:
        extract_docx(input_file, output_file)

if __name__ == "__main__":
    main()
